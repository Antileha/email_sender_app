import os
from datetime import datetime
from openpyxl import Workbook, load_workbook

HISTORY_DIR = "history"
HISTORY_FILE = os.path.join(HISTORY_DIR, "send_history.xlsx")


def ensure_history_file():
    os.makedirs(HISTORY_DIR, exist_ok=True)

    if not os.path.exists(HISTORY_FILE):
        wb = Workbook()
        ws = wb.active
        ws.title = "History"

        ws.append([
            "Дата",
            "SMTP-сервис",
            "Отправитель",
            "Получатель",
            "CC",
            "BCC",
            "Тема",
            "Вложения",
            "Статус",
            "Ошибка"
        ])

        wb.save(HISTORY_FILE)


def write_history(
    smtp_service,
    sender,
    recipient,
    cc,
    bcc,
    subject,
    attachments,
    status,
    error=""
):
    ensure_history_file()

    wb = load_workbook(HISTORY_FILE)
    ws = wb["History"]

    ws.append([
        datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        smtp_service,
        sender,
        recipient,
        cc,
        bcc,
        subject,
        attachments,
        status,
        error
    ])

    wb.save(HISTORY_FILE)